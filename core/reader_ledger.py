# -*- coding: utf-8 -*-
"""读样品台账，标定每张表的列语义，建判重索引与「位置 → 表」投票表。

台账的坑（实测）：
  1. 28 张表表头文字与内容不符的有 7 张（北部/南部山头炮孔写「样品编号/分析编号」，
     排土场/低品位堆场样/爆堆样/路边矿/倒矿样 写「送样编号/分析编号」），
     而 1-1剥离 / 前山Ⅱ号矿体-炮孔 才是真的送样在前 → **只能按内容形态认列**。
  2. 多张表是「多组并排」（剥离样 3 组、中铁十九局 3 组、原矿堆场 2 组…），
     组与组之间有偏移（原矿堆场是 A-D + F-I，E 是空列）→ 以「分析编号列」为锚点，向右取 4 列。
  3. 尾部有大量带格式空行（中铁十九局 max_row=1767 但最后数据在 1264）→
     追加位置必须用「该组最后一个分析编号行 + 1」。
"""
from __future__ import annotations

import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field

from .reader_lab import ANA_RE, normalize_ana

# 与化验结果同构的 4 列：分析编号 / 送样编号 / 品位 / 取样位置
GROUP_WIDTH = 4

STAGING_SHEET = "未知位置待定"

# 位置列特征：含中文
CN_RE = re.compile(r"[\u4e00-\u9fa5]")
# 送样编号特征：字母数字组合，基本无中文（`Ⅰ-1-BL1-7.31` 的 Ⅰ 是罗马数字，不在 CJK 区）
SAMPLE_RE = re.compile(r"^[A-Za-z0-9\u2160-\u217f\-_./ ]+$")

SKIP_SHEETS_HINT = "筛分矿样、选厂原矿数据、构造点 的编号不是 2026xxxxxx-n 形态，自动排除"


@dataclass
class ColumnRole:
    ana: int | None = None
    sample: int | None = None
    grade: int | None = None
    loc: int | None = None

    def all(self) -> list[int]:
        return [c for c in (self.ana, self.sample, self.grade, self.loc) if c is not None]


@dataclass
class Group:
    """一组并排的四列（0-based 列号）。"""
    anchor: int
    role: ColumnRole
    last_row: int = 0      # 1-based，组内最后一个「分析编号非空」的行
    rows: int = 0          # 组内有效记录数
    span: tuple[int, int] = (0, 0)


@dataclass
class SheetInfo:
    name: str
    max_row: int
    max_col: int
    groups: list[Group] = field(default_factory=list)
    is_staging: bool = False

    @property
    def ana_count(self) -> int:
        return sum(g.rows for g in self.groups)

    @property
    def usable(self) -> bool:
        return bool(self.groups) and not self.is_staging

    def anchor_cols(self) -> list[int]:
        return [g.anchor for g in self.groups]


@dataclass
class Ledger:
    path: str
    sheets: list[SheetInfo] = field(default_factory=list)
    # 判重：sheet -> 该表出现过的分析编号
    by_sheet: dict[str, set[str]] = field(default_factory=dict)
    # 反查：编号 -> 出现在哪些表
    by_ana: dict[str, set[str]] = field(default_factory=lambda: defaultdict(set))
    # 归档学习：位置文本 -> Counter(sheet)
    vote: dict[str, Counter] = field(default_factory=lambda: defaultdict(Counter))
    notes: list[str] = field(default_factory=list)

    @property
    def all_ana(self) -> set[str]:
        return set(self.by_ana)

    @property
    def total_ana(self) -> int:
        return len(self.by_ana)

    def sheet(self, name: str) -> SheetInfo | None:
        for s in self.sheets:
            if s.name == name:
                return s
        return None

    @property
    def staging_sheet(self) -> SheetInfo | None:
        return self.sheet(STAGING_SHEET)

    @property
    def usable_sheets(self) -> list[SheetInfo]:
        return [s for s in self.sheets if s.usable]


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def _is_num(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)



def _role_of_block(rows: list[tuple], anchor: int, width: int) -> ColumnRole:
    """以 anchor（0-based，分析编号列）为基准，认这一组 4 列的语义。

    ⚠ 候选列必须限制在 anchor **右侧** —— 早期版本把 anchor-1 也算进来，
    结果「剥离样」第 2/3 组的位置列被判成第 1 组的 D 列（D 列中文最多），全表位置学歪。
    anchor 左侧只在右侧找不到送样编号时才回退看（1-1剥离 / 前山Ⅱ号矿体 是送样在前）。
    """
    right = [c for c in range(anchor + 1, min(width, anchor + GROUP_WIDTH))]

    stats: dict[int, dict[str, int]] = {}

    def collect(c: int) -> dict[str, int]:
        if c in stats:
            return stats[c]
        num = cn = sample_like = total = 0
        for row in rows:
            v = row[c] if c < len(row) else None
            if v in (None, ""):
                continue
            total += 1
            if _is_num(v):
                num += 1
            s = _text(v)
            if CN_RE.search(s):
                cn += 1
            if SAMPLE_RE.match(s):
                sample_like += 1
        stats[c] = {"num": num, "cn": cn, "sample": sample_like, "total": total}
        return stats[c]

    for c in right:
        collect(c)
    if anchor - 1 >= 0:
        collect(anchor - 1)

    role = ColumnRole(ana=anchor)

    # 品位：右侧数值最多的列
    grade = max(right, key=lambda c: stats[c]["num"], default=None) if right else None
    if grade is not None and stats[grade]["num"] >= max(1, stats[grade]["total"] // 3):
        role.grade = grade
    rest = [c for c in right if c != role.grade]

    # 位置：右侧含中文最多的列（至少有 3 行中文文本，容忍位置稀疏的表）
    loc = max(rest, key=lambda c: stats[c]["cn"], default=None) if rest else None
    if loc is not None and stats[loc]["cn"] >= 3:
        role.loc = loc
    rest = [c for c in rest if c != role.loc]

    # 送样编号：右侧剩下的里最像编号的；右侧没有就回退看左侧
    sample = max(rest, key=lambda c: stats[c]["sample"], default=None) if rest else None
    if sample is not None and stats[sample]["sample"] == 0:
        sample = None
    if sample is None and anchor - 1 >= 0:
        left = stats[anchor - 1]
        if left["sample"] > 0 and left["sample"] >= left["cn"]:
            sample = anchor - 1
    role.sample = sample

    return role


def _probe_cols(ws, probe_rows: int = 300, hard_cap: int = 24) -> int:
    """探测实际用到的列数（1-based）。

    ⚠ 不能用 ws.max_column 直接 iter_rows —— 北部山头炮孔 / 南部山头炮孔的
    max_column 是 16384（整行 XFD 的格式残留），11,957 行 × 16384 列 ≈ 2 亿个单元格，
    直接把进程撑爆（实测被 OOM 杀掉、连 traceback 都来不及打）。
    """
    last = min(ws.max_row or probe_rows, probe_rows)
    top = min(ws.max_column or hard_cap, hard_cap)
    ncol = 0
    for row in ws.iter_rows(min_row=1, max_row=last, max_col=top, values_only=True):
        for i, v in enumerate(row, start=1):
            if v not in (None, ""):
                ncol = max(ncol, i)
    return max(ncol, 1)


def read_ledger(path: str) -> Ledger:
    import openpyxl

    book = openpyxl.load_workbook(path, data_only=True)
    ledger = Ledger(path=path)

    for ws in book.worksheets:
        width = _probe_cols(ws)
        if (ws.max_row or 0) == 0:
            ledger.sheets.append(SheetInfo(ws.title, 0, 0))
            continue

        rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row,
                                 max_col=width, values_only=True))
        if not rows:
            ledger.sheets.append(SheetInfo(ws.title, 0, 0))
            continue

        info = SheetInfo(
            name=ws.title,
            max_row=len(rows),
            max_col=width,
            is_staging=(ws.title == STAGING_SHEET),
        )

        # 找所有锚点：该列上「分析编号」形态命中最多的那一列（每组一个）
        hits = {}
        for c in range(width):
            n = sum(1 for row in rows if c < len(row) and isinstance(row[c], str)
                    and ANA_RE.match(normalize_ana(row[c])))
            if n:
                hits[c] = n

        # 相邻锚点间距 < 4 的视为同一组（只留命中数最多的）
        anchors: list[int] = []
        for c in sorted(hits, key=lambda c: -hits[c]):
            if all(abs(c - a) >= GROUP_WIDTH for a in anchors):
                anchors.append(c)
        anchors.sort()

        for anchor in anchors:
            role = _role_of_block(rows, anchor, width)
            grp = Group(anchor=anchor, role=role,
                        span=(max(1, anchor - 1), min(width, anchor + GROUP_WIDTH - 1)))
            for r_idx, row in enumerate(rows, start=1):
                v = row[anchor] if anchor < len(row) else None
                if isinstance(v, str) and ANA_RE.match(normalize_ana(v)):
                    grp.rows += 1
                    grp.last_row = r_idx
            info.groups.append(grp)

        ledger.sheets.append(info)

        # ---- 判重索引：全表扫描，宁可多认不漏 ----
        # 不只在锚点列里找 —— 个别表的编号会串到相邻列（例如位置列里混进一个编号），
        # 漏掉一个就会把已在台账里的记录当成"未并入"重复写一遍。
        ana_set: set[str] = set()
        for row in rows:
            for v in row:
                if isinstance(v, str):
                    a = normalize_ana(v)
                    if ANA_RE.match(a):
                        ana_set.add(a)
        ledger.by_sheet[ws.title] = ana_set
        for a in ana_set:
            ledger.by_ana[a].add(ws.title)

        # ---- 归档学习：只从锚点组的「位置列」学，暂存表不参与 ----
        if not info.is_staging:
            for grp in info.groups:
                a, l = grp.role.ana, grp.role.loc
                if a is None or l is None:
                    continue
                for row in rows:
                    v = row[a] if a < len(row) else None
                    if not (isinstance(v, str) and ANA_RE.match(normalize_ana(v))):
                        continue
                    loc = _text(row[l]) if l < len(row) else ""
                    if loc:
                        ledger.vote[loc][ws.title] += 1

    book.close()

    # 无锚点的表不可能接收数据，记一笔说明
    for s in ledger.sheets:
        if not s.groups:
            ledger.notes.append(f"「{s.name}」没有可识别的分析编号列，不参与自动并入")

    return ledger


def profile_dict(ledger: Ledger, excluded: set[str] | None = None) -> dict:
    """导出 sheet_profile.json 的内容，供人工核对。"""
    from openpyxl.utils import get_column_letter as col

    excluded = excluded or set()

    def name(idx: int | None) -> str:
        return col(idx + 1) if idx is not None else "—"

    sheets = []
    for s in ledger.sheets:
        groups = []
        for g in s.groups:
            groups.append({
                "anchor_col": name(g.anchor),
                "ana_col": name(g.role.ana),
                "sample_col": name(g.role.sample),
                "grade_col": name(g.role.grade),
                "loc_col": name(g.role.loc),
                "last_data_row": g.last_row,
                "rows": g.rows,
                "next_write_row": g.last_row + 1,
            })
        sheets.append({
            "sheet": s.name,
            "max_row": s.max_row,
            "used_cols": s.max_col,
            "is_staging": s.is_staging,
            "has_anchor": bool(s.groups),
            "writable": s.usable and s.name not in excluded,
            "excluded": s.name in excluded,
            "groups": groups,
        })
    return {
        "ledger": ledger.path,
        "total_ana": ledger.total_ana,
        "locations_learned": len(ledger.vote),
        "writable_sheets": sum(1 for s in sheets if s["writable"]),
        "sheets": sheets,
    }
