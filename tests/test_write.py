# -*- coding: utf-8 -*-
"""写入测试（计划书「难点 3」「难点 4」，验收第 5~9 条）。

只碰 `_out/test/` 下的副本，绝不动 `训练材料` 里的原件。
"""
from __future__ import annotations

import os
import unittest

import base
from core import backup, pipeline, reader_ledger, writer
from core.reader_ledger import STAGING_SHEET


class TestWriteOnCopy(unittest.TestCase):
    """整个类共享一次写入结果（写一次要 20 秒）。

    ⚠ 必须对**副本**重新跑 analyze —— 直接复用 base.get_analysis() 会把
    analysis.ledger_path 指向原件，pipeline.apply 就写到原件上去了（踩过一次）。
    """

    @classmethod
    def setUpClass(cls):
        cls.copy = base.copy_ledger("写入测试.xlsx")
        # 对副本做分析，analysis.ledger_path 才会指向副本
        cls.analysis = pipeline.analyze(cls.copy, base.LAB_DIR)
        assert os.path.abspath(cls.analysis.ledger_path) == os.path.abspath(cls.copy), \
            "analysis 必须指向测试副本，否则会写到原件"
        cls.result = pipeline.apply(cls.analysis, base.OUT_DIR, do_backup=True)
        cls.after = reader_ledger.read_ledger(cls.copy)

    @classmethod
    def tearDownClass(cls):
        # 副本留着不删（执行环境可能禁止删除文件，也无所谓，_out 本来就是临时目录）
        pass

    # ---- 数量 ----
    def test_written_total(self):
        self.assertEqual(self.result.total, 169)

    def test_per_sheet_counts(self):
        got = self.result.written
        self.assertEqual(got.get("北部山头炮孔"), 106)
        self.assertEqual(got.get("剥离样"), 24)
        self.assertEqual(got.get("中铁十九局低品位堆场样"), 24)
        self.assertEqual(got.get(STAGING_SHEET), 15)

    def test_total_ana_increases_by_158(self):
        """169 条里 11 条是跨表已存在（编号本来就在），所以总编号只增 158。"""
        before = self.analysis.ledger.total_ana
        self.assertEqual(self.after.total_ana, before + 158)

    def test_staging_sheet_created(self):
        self.assertTrue(self.result.staging_created)
        st = self.after.staging_sheet
        self.assertIsNotNone(st)
        self.assertEqual(st.ana_count, 15)

    def test_sheet_count_29(self):
        self.assertEqual(len(self.after.sheets), 29)

    # ---- 位置 ----
    def test_zhongtie_written_at_1265(self):
        grp = self.after.sheet("中铁十九局低品位堆场样").groups[0]
        self.assertEqual(grp.last_row, 1265 + 24 - 1)
        self.assertEqual(grp.last_row + 1, 1289)

    def test_no_blank_gap_before_new_rows(self):
        """新数据必须紧接原最后一行，中间不能有 500 行空白。"""
        import openpyxl
        wb = openpyxl.load_workbook(self.copy, data_only=True, read_only=True)
        ws = wb["中铁十九局低品位堆场样"]
        rows = list(ws.iter_rows(min_row=1264, max_row=1266, max_col=1, values_only=True))
        wb.close()
        self.assertTrue(all(r[0] for r in rows), "1264/1265/1266 三行都该有分析编号")

    def test_beifu_written_at_end(self):
        grp = self.after.sheet("北部山头炮孔").groups[0]
        self.assertGreater(grp.last_row, 11957)

    # ---- 内容 ----
    def test_natural_order(self):
        """新增行按 -1,-2,…,-10,-11 排，不是字典序 -1,-10,-11,…,-2。"""
        import openpyxl
        wb = openpyxl.load_workbook(self.copy, data_only=True, read_only=True)
        ws = wb["中铁十九局低品位堆场样"]
        rows = list(ws.iter_rows(min_row=1265, max_row=1274, max_col=1, values_only=True))
        wb.close()
        tail = [r[0].split("-")[-1] for r in rows]
        self.assertEqual(tail, [str(i) for i in range(1, 11)])

    def test_staging_row_content(self):
        import openpyxl
        wb = openpyxl.load_workbook(self.copy, data_only=True, read_only=True)
        ws = wb[STAGING_SHEET]
        rows = list(ws.iter_rows(min_row=1, max_row=2, max_col=6, values_only=True))
        wb.close()
        self.assertEqual(list(rows[0]), writer.STAGING_HEADER)
        self.assertEqual(rows[1][0], "2026004003-1-1")
        self.assertEqual(rows[1][1], "Y-YKDC-0927-P1")
        self.assertEqual(rows[1][4], "20260928上库选厂矿石Au -2.xls")
        self.assertEqual(rows[1][5], "原矿堆场山顶运矿样")   # 建议列要有内容

    # ---- 格式 ----
    def test_style_matches_previous_row(self):
        """比对实际样式属性（_style 里的样式表索引在保存重排后会变，不能直接比）。"""
        import openpyxl
        wb = openpyxl.load_workbook(self.copy)
        ws = wb["北部山头炮孔"]
        last = self.after.sheet("北部山头炮孔").groups[0].last_row
        for c in range(1, 5):
            new, old = ws.cell(last, c), ws.cell(last - 1, c)
            self.assertEqual(new.font.name, old.font.name, f"第 {c} 列字体不一致")
            self.assertEqual(new.font.sz, old.font.sz, f"第 {c} 列字号不一致")
            self.assertEqual(new.number_format, old.number_format, f"第 {c} 列数字格式不一致")
            self.assertEqual(new.alignment.horizontal, old.alignment.horizontal)
            self.assertEqual(bool(new.border.left.style), bool(old.border.left.style))
        self.assertEqual(ws.cell(last, 3).number_format, "0.00_ ")
        wb.close()

    def test_formulas_preserved(self):
        """data_only=False 保存时不能把公式变成值。"""
        import openpyxl
        wb = openpyxl.load_workbook(self.copy)
        self.assertEqual(wb["二堆场"]["I1282"].value, "=G1282*H1282")
        self.assertEqual(wb["多元素分析实验"]["E36"].value, "=E30/10000")
        self.assertEqual(wb["中心矿段筛分氧化矿临时堆场"]["G8"].value, "=AVERAGE(G2:G7)")
        wb.close()

    # ---- 安全 ----
    def test_backup_created(self):
        self.assertTrue(self.result.backup_path)
        self.assertTrue(os.path.exists(self.result.backup_path))
        self.assertIn(".bak", self.result.backup_path)
        self.assertGreater(os.path.getsize(self.result.backup_path), 1_000_000)

    def test_idempotent(self):
        """对已写入的副本再分析一次，应该 0 条待写。"""
        again = pipeline.analyze(self.copy, base.LAB_DIR)
        self.assertEqual(again.total, 0)


class TestBackupRotation(unittest.TestCase):
    def test_keep_limit(self):
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            folder = backup.backup_dir(tmp)
            for i in range(14):
                path = os.path.join(folder, f"x_并入0101-00{i:02d}.bak.xlsx")
                with open(path, "w") as fh:
                    fh.write("x")
                os.utime(path, (1000 + i, 1000 + i))
            removed = backup.prune(folder, keep=10)
            self.assertEqual(len(removed), 4)
            self.assertEqual(len(os.listdir(folder)), 10)


if __name__ == "__main__":
    unittest.main(verbosity=2)
